#!/usr/bin/env python3
"""Recalculate an XLSX in an isolated LibreOffice profile and report formula errors."""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from openpyxl import load_workbook


def find_soffice():
    configured = os.environ.get('SOFFICE')
    if configured:
        return configured
    found = shutil.which('soffice') or shutil.which('libreoffice')
    if found:
        return found
    mac = Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    return str(mac) if mac.exists() else None


def recalc(filename, timeout=30):
    source = Path(filename).resolve()
    if not source.is_file():
        return {'error': f'File {filename} does not exist'}
    executable = find_soffice()
    if not executable:
        return {'error': 'LibreOffice is required. Install it or set SOFFICE to its executable.'}
    try:
        with tempfile.TemporaryDirectory(prefix='agentkit-recalc-') as tmp:
            work = Path(tmp)
            output = work / 'output'
            output.mkdir()
            command = [executable, '-env:UserInstallation=' + (work / 'profile').as_uri(),
                       '--headless', '--convert-to', 'xlsx', '--outdir', str(output), str(source)]
            process = subprocess.run(command, capture_output=True, text=True, timeout=timeout)
            converted = output / source.name
            if process.returncode or not converted.is_file():
                return {'error': 'LibreOffice recalculation failed', 'details': process.stderr or process.stdout}
            # Validate before replacing the input; keep the input on conversion/read errors.
            check = load_workbook(converted, data_only=False)
            check.close()
            shutil.copyfile(converted, source)
        errors = {}
        values = load_workbook(source, data_only=True)
        for sheet in values:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'e':
                        errors.setdefault(str(cell.value), []).append(f'{sheet.title}!{cell.coordinate}')
        values.close()
        formulas = load_workbook(source, data_only=False)
        count = sum(cell.data_type == 'f' for sheet in formulas for row in sheet for cell in row)
        formulas.close()
        return {'status': 'errors_found' if errors else 'success',
                'total_errors': sum(map(len, errors.values())), 'total_formulas': count,
                'error_summary': {key: {'count': len(cells), 'locations': cells[:20]} for key, cells in errors.items()}}
    except (OSError, subprocess.TimeoutExpired, ValueError) as error:
        return {'error': str(error)}


def main():
    if len(sys.argv) < 2:
        print('Usage: recalc.py <excel_file> [timeout_seconds]', file=sys.stderr)
        return 1
    result = recalc(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 30)
    print(json.dumps(result, indent=2))
    return 1 if 'error' in result or result.get('total_errors') else 0


if __name__ == '__main__':
    sys.exit(main())
