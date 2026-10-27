#!/usr/bin/env python3
"""Create, edit, reread and render all four preserved document formats in a temp workspace."""
import argparse
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile

ROOT = Path(__file__).resolve().parent.parent


def run(script, *args, cwd):
    result = subprocess.run([sys.executable, str(ROOT / script), *map(str, args)],
                            cwd=cwd, text=True, capture_output=True, timeout=90)
    if result.returncode:
        raise RuntimeError(f'{script}: {result.stdout}\n{result.stderr}')
    return result.stdout


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default=None)
    args = parser.parse_args()
    output = Path(args.output or tempfile.mkdtemp(prefix='agentkit-documents-')).resolve()
    output.mkdir(parents=True, exist_ok=True)
    nested = output / 'nested working directory'
    nested.mkdir(exist_ok=True)
    from docx import Document
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from openpyxl import Workbook, load_workbook
    from reportlab.pdfgen import canvas
    from pypdf import PdfReader

    # DOCX: exercise existing unpack, XML edit, repack and reread helpers.
    doc = Document()
    title_style = doc.styles['Title']
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.underline = False
    for border in list(title_style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.add_heading('Document verification', 0)
    doc.add_paragraph('This sample verifies preserved document tooling and rendering.')
    doc.add_paragraph('Draft result')
    doc.save(output / 'sample.docx')
    unpacked = output / 'word-unpacked'
    run('skills/docx/ooxml/scripts/unpack.py', output / 'sample.docx', unpacked, cwd=nested)
    sys.path.insert(0, str(ROOT / 'skills/docx'))
    from scripts.document import Document as Editor
    editor = Editor(unpacked, author='AgentKit Engineer', initials='AK')
    xml = editor['word/document.xml']
    node = xml.get_node(tag='w:t', contains='Draft result')
    xml.replace_node(node, '<w:t>Verified result</w:t>')
    editor.save(validate=False)
    run('skills/docx/ooxml/scripts/pack.py', unpacked, output / 'edited.docx', cwd=nested)
    assert 'Verified result' in '\n'.join(p.text for p in Document(output / 'edited.docx').paragraphs)

    # PPTX: use original inventory and replacement scripts, then reopen native file.
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(8), Inches(1))
    title.text = 'Presentation verification'
    title.text_frame.paragraphs[0].font.size = Pt(30)
    deck.save(output / 'sample.pptx')
    inventory = output / 'inventory.json'
    run('skills/pptx/scripts/inventory.py', output / 'sample.pptx', inventory, cwd=nested)
    values = json.loads(inventory.read_text())
    first_slide = next(iter(values.values()))
    first_shape = next(iter(first_slide.values()))
    first_shape['paragraphs'] = [{'text': 'Verified presentation', 'font_size': 30}]
    replacements = output / 'replacements.json'
    replacements.write_text(json.dumps(values))
    run('skills/pptx/scripts/replace.py', output / 'sample.pptx', replacements, output / 'edited.pptx', cwd=nested)
    assert 'Verified presentation' in '\n'.join(s.text for s in Presentation(output / 'edited.pptx').slides[0].shapes if s.has_text_frame)

    # PDF: fill a native AcroForm through the preserved helper, then reread it.
    pdf = canvas.Canvas(str(output / 'sample.pdf'))
    pdf.setFont('Helvetica-Bold', 22)
    pdf.drawString(55, 770, 'PDF verification')
    pdf.setFont('Helvetica', 12)
    pdf.drawString(55, 735, 'Status')
    pdf.acroForm.textfield(name='status', x=55, y=690, width=300, height=26)
    pdf.showPage()
    pdf.save()
    fields = output / 'fields.json'
    fields.write_text(json.dumps([{'field_id': 'status', 'page': 1, 'value': 'Verified result'}]))
    run('skills/pdf/scripts/fill_fillable_fields.py', output / 'sample.pdf', fields, output / 'edited.pdf', cwd=nested)
    assert PdfReader(output / 'edited.pdf').get_fields()['status']['/V'] == 'Verified result'

    # XLSX: edit an input and use the upgraded original recalc entrypoint.
    wb = Workbook()
    ws = wb.active
    ws.title = 'Verification'
    ws.append(['Quantity', 'Unit cost', 'Total'])
    ws.append([2, 3, '=A2*B2'])
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    wb.save(output / 'sample.xlsx')
    wb = load_workbook(output / 'sample.xlsx')
    wb.active['A2'] = 4
    wb.save(output / 'edited.xlsx')
    result = json.loads(run('skills/xlsx/recalc.py', output / 'edited.xlsx', '60', cwd=nested))
    assert result['status'] == 'success' and result['total_formulas'] == 1, result
    assert load_workbook(output / 'edited.xlsx', data_only=True).active['C2'].value == 12

    rendered = {}
    for extension in ['docx', 'pptx', 'pdf', 'xlsx']:
        target = output / ('render-' + extension)
        run('skills/common/render_office.py', output / ('edited.' + extension), target, cwd=nested)
        images = list(target.glob('page-*.png'))
        assert images, extension
        rendered[extension] = len(images)
    report = {'ok': True, 'output': str(output), 'renderedPages': rendered,
              'checks': ['create', 'edit', 'reread', 'render', 'nested working directory', 'XLSX recalculation']}
    (output / 'result.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
